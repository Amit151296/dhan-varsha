"""Builds a colour-coded .xlsx workbook from an analysed fetch job, replacing
the notebook's `colorIt.to_excel(...)` step with a downloadable endpoint.
"""

from __future__ import annotations

import io
import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app import config

logger = logging.getLogger(__name__)


def build_workbook(sector_label: str, rows: list) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = (sector_label or "Sector")[:31]

    columns = list(rows[0].metrics.keys()) if rows else []
    sheet.append(["Name", *columns])

    for row in rows:
        sheet.append([row.name, *(row.metrics.get(column) for column in columns)])
        excel_row = sheet.max_row
        for col_index, column in enumerate(columns, start=2):
            highlight = row.highlights.get(column)
            if not highlight:
                continue
            cell = sheet.cell(row=excel_row, column=col_index)
            bg = highlight.background.lstrip("#")
            fg = highlight.color.lstrip("#")
            cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
            cell.font = Font(color=fg)

    for col_index in range(1, len(columns) + 2):
        sheet.column_dimensions[get_column_letter(col_index)].width = 14

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def archive_to_disk(workbook_bytes: bytes, sector_key: str, fetched_at: datetime) -> Path | None:
    """Best-effort copy to <SCREENER_EXPORT_DIR>/<year>/<MONTH>/<sector_key>.xlsx.

    Mirrors the original notebook's DATA/<year>/<month>/<sector>.xlsx layout.
    A failure here (e.g. the configured drive isn't available) is logged and
    swallowed - it must never break the actual download response.
    """
    if not config.SCREENER_EXPORT_DIR:
        return None

    directory = Path(config.SCREENER_EXPORT_DIR) / str(fetched_at.year) / fetched_at.strftime("%B").upper()
    try:
        directory.mkdir(parents=True, exist_ok=True)
        file_path = directory / f"{sector_key}.xlsx"
        file_path.write_bytes(workbook_bytes)
    except OSError:
        logger.warning("Could not archive export for '%s' to %s", sector_key, directory, exc_info=True)
        return None

    logger.info("Archived export for '%s' to %s", sector_key, file_path)
    return file_path
